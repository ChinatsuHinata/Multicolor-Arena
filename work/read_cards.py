import openpyxl
p=r'C:\Users\tzx20\Documents\test\recourse\极彩全牌表 - （圣版可检索）25-12-17.xlsx'
w=openpyxl.load_workbook(p,data_only=True,read_only=True)
for s in w:
 print('SHEET',s.title,s.max_row,s.max_column)
 for row in s.iter_rows(values_only=True):
  vals=[str(v) if v is not None else '' for v in row]
  if any(any(n in v for n in ['普通的魔法使','乐园的可爱巫女','博丽神社','极限火花','梦想封印·瞬','妖魔书','蕾米莉亚的烛台','魔理沙的元素瓶']) for v in vals): print(vals)
